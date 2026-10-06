"""Build the revised Teradata documents as PDF and editable Word (24 Sep 2026).

Same substance as the 23 Sep set, reorganized to an executive, ATS friendly
template: one column, real text, standard fonts, capitalized section headings
as on the resume, simple bullets, no tables, text boxes, images or rules.
One content model feeds both the PDF (reportlab) and the DOCX (python-docx),
so the two formats always match.
Style rule from Mustafa: no separator characters anywhere.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).parent
DATE = "24 September 2026"
ROLE = "Staff AI Engineer, Teradata"
NAME = "MUSTAFA SHOUKAT"
TAGLINE = "Senior AI Engineer, Enterprise AI Architecture and Technical Leadership"
CONTACT = [
    ("mustafashoukat.ai@gmail.com", "mailto:mustafashoukat.ai@gmail.com"),
    ("linkedin.com/in/mustafashoukat", "https://linkedin.com/in/mustafashoukat"),
    ("github.com/Mustafa-Shoukat1", "https://github.com/Mustafa-Shoukat1"),
    ("Riyadh, Saudi Arabia", None),
]
ACCENT_HEX = "1F4E79"
INK_HEX = "1A1A1A"
MUTED_HEX = "555555"

# Content model: ("h", heading) | ("p", text) | ("b", text) bullet | ("sp",) spacer.
# Inline <b>bold</b> is the only markup.

SYNOPSIS = [
    ("h", "PROFILE"),
    ("p", "Mustafa Shoukat is a Senior AI Engineer from Pakistan, based in Riyadh, with more than five years of "
          "experience taking AI systems from prototype to production. He owns the architecture of distributed, "
          "secure AI platforms end to end and has shipped more than 20 production systems for enterprise and "
          "government clients in Saudi Arabia, the UAE, the United States and the United Kingdom."),
    ("h", "CAREER TRAJECTORY"),
    ("b", "<b>Agentic AI Engineer, Watad (Navid), Riyadh.</b> Architecture owner and team lead for ByanRAG 2.2, the "
          "company's Arabic first enterprise RAG platform."),
    ("b", "<b>Independent AI Software Engineer.</b> Designs and ships production AI platforms for clients in four "
          "countries."),
    ("b", "<b>Generative AI Consultant, Aretec.</b> Production LLM and retrieval applications for business users."),
    ("b", "<b>Data Scientist, COMET Estimating.</b> ML, NLP and document processing pipelines on real world "
          "construction data."),
    ("h", "KEY ACHIEVEMENTS"),
    ("b", "<b>ByanRAG 2.2.</b> Multi tenant, Arabic first RAG platform serving two organizations and more than 150 "
          "users, deployed on premises for regulated networks, with cited answers in 2 to 3 seconds on self hosted "
          "open source models."),
    ("b", "<b>Yehia R1.</b> Designed the Claude as judge evaluation pipeline that helped take Yehia R1 to first "
          "place on the AraGen Leaderboard (0.5B to 25B)."),
    ("b", "<b>Global Holman.</b> Agentic government contracting platform that analyzes SAM.gov solicitations, "
          "supports bid decisions and drafts proposals."),
    ("b", "<b>VacantSeek 2.0.</b> Led the production hardening of a B2B real estate SaaS on AWS ECS through more "
          "than 340 merged pull requests."),
    ("b", "<b>Whoza.ai.</b> Live AI voice receptionist for UK tradespeople, built with a team of four."),
    ("h", "TECHNICAL LEADERSHIP"),
    ("p", "Leads and mentors a team of four engineers across AI, LLM, frontend, and DevOps and QA; sets the "
          "standards for testing, deployment and model integration; has trained more than 150 professionals in "
          "data science and AI engineering; and contributes to open source, including LangChain, RAGAS, "
          "Unstructured and his own tool, ytkb."),
    ("h", "CORE EXPERTISE"),
    ("p", "AI platform architecture, RAG and hybrid retrieval, agentic orchestration (LangGraph, CrewAI, MCP), "
          "distributed systems (Celery, Redis, AWS ECS), production observability, Python, Java and FastAPI."),
]

BRIEF_AREAS = [
    ("AI PLATFORM ARCHITECTURE",
     "Design and deliver highly scalable AI products built for reliability and performance.",
     "Designed ByanRAG 2.2 end to end: FastAPI services, isolated vector collections per organization, a "
     "five level access model enforced on every endpoint, full audit logging and a 14 panel admin console, "
     "packaged for on premises and air gapped deployment.",
     "Two organizations and more than 150 users share one platform with strict data isolation inside regulated "
     "networks."),
    ("RETRIEVAL AUGMENTED GENERATION",
     "Apply LLMs and retrieval to enterprise data with answers people can trust.",
     "Built hybrid retrieval that combines dense and BM25 search with Reciprocal Rank Fusion and CrossEncoder "
     "reranking, tuned for bilingual Arabic and English content and measured with Recall@5, MRR and faithfulness "
     "scoring.",
     "Source cited answers in 2 to 3 seconds on self hosted open source models."),
    ("ORCHESTRATION AND AGENTIC AI",
     "Deliver agentic AI and large language model systems.",
     "Led multi agent workflows in LangGraph and CrewAI with tool calling and multi step planning: Global Holman "
     "agents read SAM.gov solicitations, analyze requirements and draft proposals, and the Whoza.ai voice agent "
     "answers and qualifies calls. Built my own MCP server with Arabic tools and use MCP tools in daily "
     "development.",
     "Analyst work that took days is now automated on a provider agnostic LLM layer."),
    ("DISTRIBUTED SYSTEMS",
     "Apply distributed system patterns with strong data consistency.",
     "Use Celery and Redis for background ingestion and agent jobs, async PostgreSQL, streaming over WebSockets "
     "and server sent events, and idempotent, signature verified webhooks. Moved VacantSeek 2.0 onto AWS ECS "
     "with CodeBuild pipelines and S3 storage.",
     "Work continues when a worker fails, and retries never create duplicate records."),
    ("PRODUCTION READINESS",
     "Resolve performance bottlenecks and production issues and maintain code quality.",
     "Run services in Docker behind nginx with TLS and CI/CD, with more than 600 tests on ByanRAG, monitoring in "
     "Prometheus, Grafana and Loki, tracing in LangSmith and Sentry, and guardrails against prompt injection. Led "
     "more than 340 merged pull requests of hardening on VacantSeek 2.0.",
     "Issues surface on dashboards before users report them; the same evaluation discipline helped take Yehia R1 "
     "to first place on AraGen."),
    ("TECHNICAL LEADERSHIP",
     "Mentor engineers and influence architecture decisions.",
     "Lead and mentor the ByanRAG team of four, own its architecture decisions and set its standards for testing, "
     "deployment and code review. Trained more than 150 professionals and work directly with clients in four "
     "countries.",
     "The team ships to one shared standard."),
]


def brief_blocks():
    blocks = [("p", "This brief maps my experience to the core requirements of the Staff AI Engineer role. Each "
                    "area gives the requirement, what I delivered and the result.")]
    for title, req, did, result in BRIEF_AREAS:
        blocks += [("h", title),
                   ("b", f"<b>Requirement.</b> {req}"),
                   ("b", f"<b>Delivered.</b> {did}"),
                   ("b", f"<b>Result.</b> {result}")]
    blocks += [
        ("h", "ENGINEERING FOUNDATIONS"),
        ("p", "Python (primary), Java, TypeScript and SQL; RESTful APIs in FastAPI; PostgreSQL, MongoDB, Redis and "
              "pgvector; AWS, GCP and Azure; Docker and AWS ECS; GitHub Copilot and MCP based coding assistants. "
              "My orchestration and messaging patterns from Docker, ECS, Celery and Redis carry directly to "
              "Kubernetes and Kafka."),
    ]
    return blocks


PURPOSE = [
    ("p", "Dear Teradata Hiring Team,"),
    ("p", "I am applying for the Staff AI Engineer role because it sits where I have chosen to build my career: "
          "trustworthy, production grade AI on top of enterprise data. Teradata's Enterprise Vector Store and "
          "Enterprise AgentStack show a clear direction, bringing retrieval and agents to governed data at scale. "
          "That is the problem I have spent the last several years solving in production."),
    ("p", "My strongest preparation is ByanRAG 2.2, the Arabic first RAG platform I designed for regulated "
          "organizations in Saudi Arabia and now lead with a team of four. It isolates each organization's data, "
          "enforces access on every request, records every action and returns cited answers in 2 to 3 seconds on "
          "self hosted open source models, today for two organizations and more than 150 users. Alongside it, I "
          "built the evaluation pipeline that helped take Yehia R1 to first place on the AraGen Leaderboard. Both "
          "taught me that enterprise AI is judged on reliability and evidence, not on a demo."),
    ("p", "In my first months I would learn Teradata's platform, codebase and customers in depth, contribute "
          "working code early, and help set clear baselines for evaluation, observability and testing across the "
          "AI products. I would also support the engineers around me through design reviews, code reviews and "
          "mentoring, as I do with my current team."),
    ("p", "At this stage of my career I want to grow as a Staff engineer: shaping architecture across teams, "
          "raising engineering standards and building AI products that enterprises trust. I am based in Riyadh and "
          "flexible on the working arrangement that suits the team. Thank you for considering my application."),
    ("sp",),
    ("p", "Sincerely,"),
    ("p", "<b>Mustafa Shoukat</b>"),
]

DOCS = [
    ("Mustafa_Shoukat_Executive_Synopsis", "Executive Synopsis", SYNOPSIS),
    ("Mustafa_Shoukat_Executive_Positioning_Brief", "Executive Positioning Brief", brief_blocks()),
    ("Mustafa_Shoukat_Statement_of_Purpose", "Statement of Purpose", PURPOSE),
]


# ------------------------------------------------------------------ PDF
PS = {
    "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=16, leading=19,
                           textColor=HexColor("#000000"), alignment=TA_CENTER),
    "tag": ParagraphStyle("tag", fontName="Helvetica", fontSize=9.5, leading=12,
                          textColor=HexColor("#" + ACCENT_HEX), alignment=TA_CENTER, spaceBefore=2),
    "contact": ParagraphStyle("contact", fontName="Helvetica", fontSize=8.8, leading=11.5,
                              textColor=HexColor("#" + MUTED_HEX), alignment=TA_CENTER, spaceBefore=2),
    "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=13, leading=16,
                            textColor=HexColor("#" + ACCENT_HEX), alignment=TA_LEFT, spaceBefore=14),
    "meta": ParagraphStyle("meta", fontName="Helvetica", fontSize=9, leading=12,
                           textColor=HexColor("#" + MUTED_HEX), spaceAfter=4),
    "h": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=10, leading=12.5,
                        textColor=HexColor("#" + ACCENT_HEX), spaceBefore=9, spaceAfter=3),
    "p": ParagraphStyle("p", fontName="Helvetica", fontSize=10, leading=13.6,
                        textColor=HexColor("#" + INK_HEX), alignment=TA_JUSTIFY, spaceAfter=6),
    "b": ParagraphStyle("b", fontName="Helvetica", fontSize=10, leading=13.4,
                        textColor=HexColor("#" + INK_HEX), alignment=TA_JUSTIFY,
                        leftIndent=14, bulletIndent=3, spaceAfter=3),
}


def pdf_header(story, title):
    gap = "&nbsp;" * 4
    parts = [f'<a href="{url}" color="#1155CC">{text}</a>' if url else text for text, url in CONTACT]
    story += [Paragraph(NAME, PS["name"]), Paragraph(TAGLINE, PS["tag"]),
              Paragraph(gap.join(parts), PS["contact"]),
              Paragraph(title, PS["title"]), Paragraph(f"{ROLE}, {DATE}", PS["meta"])]


def build_pdf(stem, title, blocks):
    doc = SimpleDocTemplate(str(OUT / f"{stem}.pdf"), pagesize=letter,
                            leftMargin=0.85 * inch, rightMargin=0.85 * inch,
                            topMargin=0.6 * inch, bottomMargin=0.6 * inch,
                            title=f"Mustafa Shoukat, {title}", author="Mustafa Shoukat",
                            subject=f"{title} for {ROLE}")
    story = []
    pdf_header(story, title)
    section = []  # each heading and its blocks stay on one page

    def flush():
        if section:
            story.append(KeepTogether(list(section)))
            section.clear()

    for block in blocks:
        kind = block[0]
        if kind == "h":
            flush()
            section.append(Paragraph(block[1], PS["h"]))
        elif kind == "sp":
            section.append(Spacer(1, 6))
        elif kind == "b":
            section.append(Paragraph(block[1], PS["b"], bulletText="•"))
        else:
            section.append(Paragraph(block[1], PS["p"]))
    flush()
    doc.build(story)


# ----------------------------------------------------------------- DOCX
def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def add_runs(par, text, size, color=INK_HEX):
    for i, chunk in enumerate(re.split(r"</?b>", text)):
        if chunk:
            run = par.add_run(chunk)
            run.bold = i % 2 == 1
            run.font.size = Pt(size)
            run.font.name = "Arial"
            run.font.color.rgb = rgb(color)


def spacing(par, before=0, after=0, line=1.15):
    fmt = par.paragraph_format
    fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(before), Pt(after), line


def build_docx(stem, title, blocks):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin = sec.bottom_margin = Inches(0.6)
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    d.core_properties.title = f"Mustafa Shoukat, {title}"
    d.core_properties.author = "Mustafa Shoukat"

    for text, size, color, bold, after in [
        (NAME, 16, "000000", True, 1),
        (TAGLINE, 9.5, ACCENT_HEX, False, 1),
        ("    ".join(t for t, _ in CONTACT), 8.8, MUTED_HEX, False, 10),
    ]:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, f"<b>{text}</b>" if bold else text, size, color)
        spacing(p, after=after)
    p = d.add_paragraph()
    add_runs(p, f"<b>{title}</b>", 13, ACCENT_HEX)
    spacing(p, before=4, after=1)
    p = d.add_paragraph()
    add_runs(p, f"{ROLE}, {DATE}", 9, MUTED_HEX)
    spacing(p, after=4)

    for block in blocks:
        kind = block[0]
        if kind == "sp":
            spacing(d.add_paragraph(), after=2)
        elif kind == "h":
            p = d.add_paragraph()
            add_runs(p, f"<b>{block[1]}</b>", 10, ACCENT_HEX)
            spacing(p, before=8, after=2)
            p.paragraph_format.keep_with_next = True
        elif kind == "b":
            p = d.add_paragraph(style="List Bullet")
            add_runs(p, block[1], 10)
            spacing(p, after=2)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            p = d.add_paragraph()
            add_runs(p, block[1], 10)
            spacing(p, after=5)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.save(str(OUT / f"{stem}.docx"))


if __name__ == "__main__":
    for stem, title, blocks in DOCS:
        build_pdf(stem, title, blocks)
        build_docx(stem, title, blocks)
        print("Wrote", stem, ".pdf and .docx")
