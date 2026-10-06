"""Build the MetLife Japan CV as PDF and editable Word (28 Sep 2026).

Tailored to the AI & Prompt/Context Engineer JD sent by Hays Japan: RAG pipelines,
evaluation hooks, hallucination reduction, prompt and tool design, APIs, data work
and direct work with business owners. Same template as the 24 Sep Teradata set:
one column, real text, standard fonts, no tables, text boxes, images or rules.
One content model feeds both the PDF (reportlab) and the DOCX (python-docx).
Style rule from Mustafa: no separator characters anywhere.

Facts differ from the 19 Sep Teradata resume on purpose, because a Japanese work
visa and MetLife's background check both verify them:
  education is "in progress, expected 2027" (live LinkedIn: Sep 2023 to Aug 2027),
  no MS NUST line, open source PRs described as submitted, "5+ years", no "air gapped".
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate

OUT = Path(__file__).parent
STEM = "Mustafa_Shoukat_CV_MetLife_Japan"
NAME = "MUSTAFA SHOUKAT"
TAGLINE = "AI Engineer, Production RAG, Agent Workflows and LLM Evaluation"
CONTACT = [
    ("mustafashoukat.ai@gmail.com", "mailto:mustafashoukat.ai@gmail.com"),
    ("linkedin.com/in/mustafashoukat", "https://linkedin.com/in/mustafashoukat"),
    ("github.com/Mustafa-Shoukat1", "https://github.com/Mustafa-Shoukat1"),
]
LOCATION = "Based in Riyadh, Saudi Arabia. Open to relocating to Tokyo."
ACCENT_HEX = "1F4E79"
INK_HEX = "1A1A1A"
MUTED_HEX = "555555"

# Content model: ("h", heading) | ("p", text) | ("b", text) bullet
# | ("r", role, place and dates) | ("s", label, text) skill line.
# Inline <b>bold</b> is the only markup.

CV = [
    ("h", "PROFILE"),
    ("p", "AI engineer with more than five years of experience taking RAG and agent systems into production for "
          "enterprise and government clients in Saudi Arabia, the UAE, the UK and the US. I lead the retrieval and "
          "evaluation work on a multi tenant enterprise RAG platform whose answers cite their sources, are scored "
          "for faithfulness against those sources, and are measured for quality, latency and cost before every "
          "release. I work directly with business owners to turn their problems into working AI applications, and "
          "I have solved the problems that non English enterprise documents bring to retrieval: Arabic first "
          "search, native language test sets, and multilingual (Arabic and CJK) support contributed to Ragas."),

    ("h", "CORE SKILLS"),
    ("s", "RAG pipelines", "Document ingestion with Docling and OCR, token aware chunking with Unicode "
          "normalization, content hash deduplication, metadata and access scope filters, hybrid dense and BM25 "
          "retrieval with Reciprocal Rank Fusion, cross encoder reranking, query reformulation, cited answers. "
          "Qdrant, Milvus, FAISS, ChromaDB, pgvector."),
    ("s", "Prompt and context engineering", "System instructions, source usage rules, structured outputs, tool "
          "definitions, LLM as judge rubrics, context selection within token budgets, response caching."),
    ("s", "Evaluation and hallucination reduction", "NLI faithfulness scoring, golden test sets, Recall@k and MRR, "
          "Ragas, LangSmith tracing, regression tests and release gates in CI."),
    ("s", "Agents", "LangGraph, LangChain, CrewAI, AutoGen, LlamaIndex, MCP (own server with Arabic tools), tool "
          "calling, multi step workflows with human handover."),
    ("s", "Backend and APIs", "Python, FastAPI, Flask, Django, Node.js with Express, TypeScript, REST, SSE and "
          "WebSockets, Celery and Redis, JWT and OAuth2, Casbin RBAC, rate limiting, retries, audit logging."),
    ("s", "Data", "Pandas, NumPy, PostgreSQL, MongoDB, Redis, SQLAlchemy; statistics for data quality and usage "
          "analysis."),
    ("s", "LLM services", "OpenAI API, Gemini, Hugging Face Inference, Llama and other open models, Ollama."),
    ("s", "Cloud and delivery", "Docker, GitHub Actions CI/CD, AWS (ECS, S3, App Runner), GCP (Cloud Run), Azure "
          "(Blob Storage), nginx, Prometheus, Grafana, Loki, Sentry; pull request based delivery with code review; "
          "Claude Code and GitHub Copilot in daily work."),

    ("h", "EXPERIENCE"),
    ("r", "AI Engineer, Navid", "Riyadh, Saudi Arabia. Jun 2026 to Present"),
    ("b", "Lead the retrieval and evaluation workstream of Navid's Arabic first enterprise RAG platform with a team of "
          "four: AI, LLM, frontend, and DevOps and QA engineers."),
    ("b", "Rebuilt the evaluation harness around a native Arabic golden set, so retrieval quality, answer "
          "faithfulness, latency and cost per query are measured before every release instead of by spot checks."),
    ("b", "Set the release standards the platform ships against: coverage gate in CI, security scanning, "
          "reproducible container builds and audit logging on every privileged action."),

    ("r", "AI Software Engineer, Watad (Navid group)", "Riyadh, Saudi Arabia. Nov 2023 to Jun 2026"),
    ("b", "<b>Lead engineer on ByanRAG</b>, a multi tenant Arabic and English document intelligence platform for "
          "regulated organizations. In production for two organizations and more than 150 users, returning cited "
          "answers in 2 to 3 seconds on open source models."),
    ("b", "<b>Built the pipeline end to end:</b> Docling conversion of 13+ file types with OCR, token aware chunking "
          "(64 to 4096 tokens) with Unicode normalization and content hash deduplication, dense and BM25 hybrid "
          "search fused with Reciprocal Rank Fusion, multilingual cross encoder reranking and query reformulation."),
    ("b", "<b>Made unsupported answers visible:</b> NLI based faithfulness scoring checks each answer against its "
          "source passages and shows the confidence to the user next to the citations."),
    ("b", "<b>Grounded answers in the right sources:</b> a separate vector collection per organization; document "
          "scopes (organization, department, private) enforced as metadata filters on every query; five level RBAC "
          "(Casbin, 32 policy rules) on every API endpoint."),
    ("b", "<b>Secured the APIs for regulated networks:</b> short lived JWT with refresh rotation and revocation, "
          "account lockout, rate limits on 17 endpoints, input validation, an upload whitelist, encrypted "
          "credential storage and an exportable audit log."),
    ("b", "<b>Instrumented production:</b> 12+ Prometheus metrics (RAG query latency, LLM request duration, cache "
          "hit rate, active streams), Loki logs and Grafana dashboards, plus a Redis response cache keyed by query "
          "and access scope that avoids repeat LLM calls."),
    ("b", "<b>Held quality in CI:</b> 617 backend and 217 frontend tests, a 65% coverage gate, Bandit scans, "
          "Playwright end to end tests and Locust load tests, all enforced in GitHub Actions."),
    ("b", "<b>Designed the LLM as judge reward pipeline</b> (3C3H rubric) used to train Navid's Yehia 7B model, which "
          "ranked first on the AraGen leaderboard for 0.5B to 25B models in February 2025."),

    ("r", "Independent AI Engineer, Contract", "Remote, clients in Saudi Arabia, the UAE, the UK and the US. "
          "2024 to Present"),
    ("b", "Gather requirements directly with business owners, propose the solution and deliver it to production."),
    ("b", "<b>VacantSeek</b>, B2B real estate SaaS: primary engineer, more than 800 commits and 340 merged pull "
          "requests in 2026, including production hardening and the move to AWS ECS."),
    ("b", "<b>Global Holman</b>, government contracting: LangGraph agents with tool calling that pull SAM.gov "
          "solicitations, analyze requirements over a RAG index, support bid decisions and draft proposal PDFs, on "
          "a provider agnostic LLM layer."),
    ("b", "<b>Whoza.ai</b>: AI voice receptionist for UK tradespeople, built with a team of four (FastAPI, Twilio, "
          "Gemini, React)."),
    ("b", "<b>Property management CRM</b>: Node.js and Express with PostgreSQL, OpenAI lead scoring and email "
          "drafting, RBAC and audit middleware."),

    ("r", "Generative AI Consultant, Aretec", "Lahore, Pakistan (remote contract). Jul 2023 to Aug 2024"),
    ("b", "Built LLM applications for business users with LangChain and LangGraph: document question answering, "
          "recruitment automation (sourcing, resume screening, interview scheduling) and structured extraction."),
    ("b", "Shipped AutoRec, a scraping and extraction SaaS: FastAPI with Celery and Redis workers, WebSocket job "
          "progress, Stripe billing and a Next.js frontend."),

    ("r", "Data Scientist, COMET Estimating LLC", "Remote (US). Jul 2021 to Jun 2023"),
    ("b", "Built Python data pipelines with Pandas and NumPy for construction estimating data: cleaning inconsistent "
          "formats and missing fields, validation checks and QA workflows, deployed on Google Cloud Run with CI/CD."),
    ("b", "Built predictive models and a lead scoring system with analyst review of automated decisions."),

    ("h", "OPEN SOURCE AND TEACHING"),
    ("b", "<b>Ragas:</b> three pull requests submitted adding Arabic and CJK multilingual support, import aliases and "
          "fixes, with 62 tests. <b>Unstructured:</b> PaddleOCR language mapping fix with tests. <b>OpenWA:</b> "
          "merged contribution."),
    ("b", "<b>ytkb</b> (MIT): command line tool that turns video into a searchable multilingual knowledge base, "
          "tested in six languages."),
    ("b", "<b>Arabic RAG Bench</b>: reinforcement learning environment that plants retrieval bugs for agents to "
          "diagnose, with deterministic Recall@5 scoring and 24 tests."),
    ("b", "Trained more than 150 professionals in data science and AI engineering."),

    ("h", "EDUCATION"),
    ("p", "<b>BS Computer Science</b>, Virtual University of Pakistan. In progress, expected 2027."),
]


# ------------------------------------------------------------------ PDF
PS = {
    "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=16, leading=19,
                           textColor=HexColor("#000000"), alignment=TA_CENTER),
    "tag": ParagraphStyle("tag", fontName="Helvetica", fontSize=10, leading=12.5,
                          textColor=HexColor("#" + ACCENT_HEX), alignment=TA_CENTER, spaceBefore=2),
    "contact": ParagraphStyle("contact", fontName="Helvetica", fontSize=8.8, leading=11.5,
                              textColor=HexColor("#" + MUTED_HEX), alignment=TA_CENTER, spaceBefore=2,
                              spaceAfter=2),
    "h": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=10.5, leading=13,
                        textColor=HexColor("#" + ACCENT_HEX), spaceBefore=9, spaceAfter=3),
    "p": ParagraphStyle("p", fontName="Helvetica", fontSize=9.4, leading=12.4,
                        textColor=HexColor("#" + INK_HEX), alignment=TA_JUSTIFY, spaceAfter=3),
    "s": ParagraphStyle("s", fontName="Helvetica", fontSize=9.2, leading=12,
                        textColor=HexColor("#" + INK_HEX), spaceAfter=2),
    "role": ParagraphStyle("role", fontName="Helvetica-Bold", fontSize=9.8, leading=12,
                           textColor=HexColor("#000000"), spaceBefore=6),
    "meta": ParagraphStyle("meta", fontName="Helvetica", fontSize=8.8, leading=11,
                           textColor=HexColor("#" + MUTED_HEX), spaceAfter=2),
    "b": ParagraphStyle("b", fontName="Helvetica", fontSize=9.4, leading=12.2,
                        textColor=HexColor("#" + INK_HEX), alignment=TA_JUSTIFY,
                        leftIndent=13, bulletIndent=3, spaceAfter=2),
}


def build_pdf():
    doc = SimpleDocTemplate(str(OUT / f"{STEM}.pdf"), pagesize=A4,
                            leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                            topMargin=1.4 * cm, bottomMargin=1.4 * cm,
                            title="Mustafa Shoukat, CV", author="Mustafa Shoukat",
                            subject="CV for AI and Prompt and Context Engineer, MetLife Japan")
    gap = "&nbsp;" * 4
    parts = [f'<a href="{url}" color="#1155CC">{text}</a>' if url else text for text, url in CONTACT]
    story = [Paragraph(NAME, PS["name"]), Paragraph(TAGLINE, PS["tag"]),
             Paragraph(gap.join(parts), PS["contact"]), Paragraph(LOCATION, PS["contact"])]
    group = []  # a heading or role stays on the page with its first item only

    for block in CV:
        kind = block[0]
        if kind == "h":
            group.append(Paragraph(block[1], PS["h"]))
            continue
        if kind == "r":
            group += [Paragraph(block[1], PS["role"]), Paragraph(block[2], PS["meta"])]
            continue
        if kind == "s":
            group.append(Paragraph(f"<b>{block[1]}:</b> {block[2]}", PS["s"]))
        elif kind == "b":
            group.append(Paragraph(block[1], PS["b"], bulletText="•"))
        else:
            group.append(Paragraph(block[1], PS["p"]))
        story.append(KeepTogether(list(group)))
        group.clear()
    doc.build(story)


# ----------------------------------------------------------------- DOCX
def add_runs(par, text, size, color=INK_HEX):
    for i, chunk in enumerate(re.split(r"</?b>", text)):
        if chunk:
            run = par.add_run(chunk)
            run.bold = i % 2 == 1
            run.font.size = Pt(size)
            run.font.name = "Arial"
            run.font.color.rgb = RGBColor.from_string(color)


def spacing(par, before=0, after=0, line=1.1):
    fmt = par.paragraph_format
    fmt.space_before, fmt.space_after, fmt.line_spacing = Pt(before), Pt(after), line


def build_docx():
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    d.styles["Normal"].font.name = "Arial"
    d.styles["Normal"].font.size = Pt(9.5)
    d.core_properties.title = "Mustafa Shoukat, CV"
    d.core_properties.author = "Mustafa Shoukat"

    for text, size, color, bold, after in [
        (NAME, 16, "000000", True, 1),
        (TAGLINE, 10, ACCENT_HEX, False, 1),
        ("    ".join(t for t, _ in CONTACT), 8.8, MUTED_HEX, False, 1),
        (LOCATION, 8.8, MUTED_HEX, False, 6),
    ]:
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, f"<b>{text}</b>" if bold else text, size, color)
        spacing(p, after=after)

    for block in CV:
        kind = block[0]
        if kind == "h":
            p = d.add_paragraph()
            add_runs(p, f"<b>{block[1]}</b>", 10.5, ACCENT_HEX)
            spacing(p, before=8, after=2)
            p.paragraph_format.keep_with_next = True
        elif kind == "r":
            p = d.add_paragraph()
            add_runs(p, f"<b>{block[1]}</b>", 9.8, "000000")
            spacing(p, before=5)
            p.paragraph_format.keep_with_next = True
            p = d.add_paragraph()
            add_runs(p, block[2], 8.8, MUTED_HEX)
            spacing(p, after=2)
            p.paragraph_format.keep_with_next = True
        elif kind == "s":
            p = d.add_paragraph()
            add_runs(p, f"<b>{block[1]}:</b> {block[2]}", 9.2)
            spacing(p, after=2)
        elif kind == "b":
            p = d.add_paragraph(style="List Bullet")
            add_runs(p, block[1], 9.4)
            spacing(p, after=2)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            p = d.add_paragraph()
            add_runs(p, block[1], 9.4)
            spacing(p, after=3)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    d.save(str(OUT / f"{STEM}.docx"))


if __name__ == "__main__":
    build_pdf()
    build_docx()
    print("Wrote", STEM, ".pdf and .docx")
